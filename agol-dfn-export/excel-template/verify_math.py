#!/usr/bin/env python3
"""
Independent check of the Schedule tab's arithmetic.

Reimplements in Python exactly what the workbook's formulas compute -- the same
flat-earth distance, the same argmin, the same distance gate -- against the same
sample data, so the expected answers are derived twice by different means.

    python3 verify_math.py
"""

import math
import sys

from openpyxl import load_workbook

FT_PER_DEG = 364000.0
MAX_FT = 500.0

SPLITTERS = [
    ("SPL-001", "Primary 1x8", "1x8", -83.000000, 40.000000),
    ("SPL-002", "Secondary 1x16", None, -83.001000, 40.000500),
    ("SPL-003", 7, None, -83.002000, 40.001000),
    ("SPL-004", None, None, -83.050000, 40.050000),
]
ADDRESSES = [
    ("100", "N HIGH ST", -83.000050, 40.000050),
    ("212", "W 5TH AVE", -83.001050, 40.000550),
    ("415", "E LANE AVE", -83.002050, 40.001050),
]
# Substring rules, mirroring the workbook's SEARCH-based lookup.
TIER_RULES = [
    # (value or word, tier, ratio supplied here or None)
    ("Primary", "Primary", None),
    ("Secondary", "Secondary", None),
    ("Tertiary", "Tertiary", None),
    (7, "Tertiary", "1x128"),
]


def matched_rule(raw):
    """Index of the rule the workbook's M column lands on; None for no match."""
    if raw is None or raw == "":
        return None
    # Exact first, compared as text so 7 and "7" are the same value.
    for i, (needle, _t, _r) in enumerate(TIER_RULES):
        if str(needle) == str(raw):
            return i
    # Then substring, last match winning as MAX does.
    hits = [i for i, (needle, _t, _r) in enumerate(TIER_RULES)
            if needle != "" and str(needle).lower() in str(raw).lower()]
    return hits[-1] if hits else None


def resolve_tier(raw):
    """What the Tier column computes."""
    i = matched_rule(raw)
    return TIER_RULES[i][1] if i is not None else ""


def resolve_ratio(raw):
    """What the Ratio column computes with no separate ratio field named."""
    i = matched_rule(raw)
    if i is not None and TIER_RULES[i][2]:
        return TIER_RULES[i][2]          # ratio typed beside the matched rule
    if raw is None:
        return ""
    text = str(raw)
    at = text.lower().find("1x")
    return text[at:at + 20].strip() if at >= 0 else ""

EXPECTED = {
    # splitter: (tier, nearest address, ratio)
    "SPL-001": ("Primary", "100 N HIGH ST", "1x8"),
    "SPL-002": ("Secondary", "212 W 5TH AVE", "1x16"),
    "SPL-003": ("Tertiary", "415 E LANE AVE", "1x128"),
    "SPL-004": ("", "", ""),  # no tier value at all, and no address within the gate
}


def dist2_ft(slon, slat, alon, alat):
    """The exact expression the K column evaluates, in square feet."""
    dx = (alon - slon) * math.cos(math.radians(slat)) * FT_PER_DEG
    dy = (alat - slat) * FT_PER_DEG
    return dx * dx + dy * dy


def main() -> int:
    failures = []
    print(f"{'splitter':<10} {'tier':<10} {'ratio':<7} {'nearest':<16} {'dist ft':>9}  verdict")
    print("-" * 68)

    for name, raw_tier, _ratio, lon, lat in SPLITTERS:
        tier = resolve_tier(raw_tier)
        ratio = resolve_ratio(raw_tier)
        d2 = [dist2_ft(lon, lat, alon, alat) for _n, _s, alon, alat in ADDRESSES]
        best = min(d2)
        idx = d2.index(best)
        dist = math.sqrt(best)
        address = f"{ADDRESSES[idx][0]} {ADDRESSES[idx][1]}" if dist <= MAX_FT else ""

        want_tier, want_addr, want_ratio = EXPECTED[name]
        ok = (tier == want_tier) and (address == want_addr) and (ratio == want_ratio)
        if not ok:
            failures.append(
                f"{name}: got ({tier!r},{address!r},{ratio!r}) "
                f"want ({want_tier!r},{want_addr!r},{want_ratio!r})"
            )
        print(f"{name:<10} {tier or '-':<10} {ratio or '-':<7} {address or '-':<16} "
              f"{dist:>9.1f}  {'ok' if ok else 'MISMATCH'}")

    # The gate must actually be doing something: SPL-004 has to be beyond it.
    lon, lat = SPLITTERS[3][3], SPLITTERS[3][4]
    far = math.sqrt(min(dist2_ft(lon, lat, a[2], a[3]) for a in ADDRESSES))
    print(f"\nSPL-004 nearest address is {far:,.0f} ft away (limit {MAX_FT:,.0f}) -> correctly blanked")
    if far <= MAX_FT:
        failures.append("SPL-004 is not beyond the distance gate; the sample no longer tests it")

    # Cross-check against haversine. Both are approximations of an ellipsoid, so
    # they will not agree exactly: the workbook's 364,000 ft/degree is a mid-latitude
    # WGS84 value (364,272 at 40N) while haversine here uses a spherical earth
    # (364,813). That is a ~0.22% scale difference applied uniformly, so what has to
    # hold is that the CHOICE of nearest address is identical -- a uniform scaling
    # cannot change an argmin -- and that reported distances stay within 0.3%.
    worst_pct = 0.0
    for name, _t, _r, slon, slat in SPLITTERS:
        flat = [math.sqrt(dist2_ft(slon, slat, a[2], a[3])) for a in ADDRESSES]
        hav = []
        for _num, _st, alon, alat in ADDRESSES:
            p1, p2 = math.radians(slat), math.radians(alat)
            a = (math.sin((p2 - p1) / 2) ** 2
                 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(alon - slon) / 2) ** 2)
            hav.append(2 * 20902231.0 * math.asin(math.sqrt(a)))
        if flat.index(min(flat)) != hav.index(min(hav)):
            failures.append(f"{name}: flat-earth and haversine pick different nearest addresses")
        for f, h in zip(flat, hav):
            worst_pct = max(worst_pct, abs(f - h) / h * 100)
    print(f"nearest-address choice matches haversine for all {len(SPLITTERS)} splitters")
    print(f"distance scale vs haversine: worst {worst_pct:.3f}% (uniform, cannot reorder)")
    if worst_pct > 0.3:
        failures.append(f"distance scale differs from haversine by {worst_pct:.2f}%")

    # The workbook must still contain formulas, not baked-in answers.
    wb = load_workbook("DFN-Splitter-Matcher.xlsx")
    ws = wb["Schedule"]
    for cell in ("B6", "D6", "H6", "I6", "K6"):
        if not str(ws[cell].value).startswith("="):
            failures.append(f"Schedule!{cell} is not a formula")
    if any("INDIRECT" in str(ws[f"{c}6"].value) for c in "ABCDEFGHIJKL"):
        failures.append("a volatile INDIRECT survived in the Schedule row")
    print("workbook still formula-driven, no volatile INDIRECT")

    if failures:
        print("\nFAILURES:")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("\nall math checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
