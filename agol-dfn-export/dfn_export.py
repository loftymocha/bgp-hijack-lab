#!/usr/bin/env python3
"""
dfn_export.py -- pull every layer out of an ArcGIS Online (AGOL) web map or
feature service and write it to a formatted Excel workbook.

Built for DFN splitter testing: alongside the raw layer dump it produces a
"Splitters" worksheet with the tier label (Primary / Secondary / Tertiary) and
the nearest street address already filled in for each splitter.

Three modes:

  discover  print every layer in the map, with field names and types.
            Run this first -- you need it to fill in config.yaml.

  profile   print the distinct values of the fields you are thinking of using
            for the tier label, so you can write accurate keyword rules.

  export    do the real work: query all features, classify tiers, resolve
            addresses, write the workbook.

No ArcGIS Python API required -- this talks to the plain REST endpoints, so it
installs with pip on a locked-down work laptop.
"""

from __future__ import annotations

import argparse
import csv
import getpass
import json
import math
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

import requests

try:
    import yaml
except ImportError:  # pragma: no cover - guarded at startup
    yaml = None

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# --------------------------------------------------------------------------
# defaults
# --------------------------------------------------------------------------

DEFAULT_CONFIG: dict[str, Any] = {
    "portal": "https://www.arcgis.com",
    "auth": {"mode": "prompt", "username": None, "api_key_env": "AGOL_API_KEY"},
    "source": {"webmap_item_id": None, "service_urls": []},
    "query": {
        "where": "1=1",
        "out_sr": 4326,
        "page_size": 1000,
        "max_features_per_layer": 100000,
        "decode_domains": True,
        "skip_layers": [],
        "only_layers": [],
    },
    "splitters": {
        "layers": ["splitter"],
        "name_fields": ["SPLITTER_ID", "SPLTR_ID", "NAME", "LABEL", "FACILITY_ID"],
        "ratio_fields": ["SPLIT_RATIO", "RATIO", "PORTS", "PORT_COUNT"],
        "tier": {
            "fields": ["TIER", "LEVEL", "SPLITTER_LEVEL", "SPLITTER_TYPE", "TYPE"],
            "keywords": {
                "Primary": ["primary", "prim", "p1", "level 1", "l1", "tier 1", "t1", "1st"],
                "Secondary": ["secondary", "sec", "p2", "level 2", "l2", "tier 2", "t2", "2nd"],
                "Tertiary": ["tertiary", "tert", "p3", "level 3", "l3", "tier 3", "t3", "3rd"],
            },
            "ratio_map": {},
            "name_patterns": {},
        },
    },
    "address": {
        "source": "layer",  # layer | geocode | none
        "layer": "address",
        "fields": ["FULL_ADDRESS"],
        "max_distance_ft": 500,
        "geocode_url": (
            "https://geocode-api.arcgis.com/arcgis/rest/services/"
            "World/GeocodeServer/reverseGeocode"
        ),
        "geocode_cache": "geocode_cache.json",
    },
    "output": {
        "workbook": "dfn_export.xlsx",
        "csv_dir": None,
        "max_column_width": 55,
    },
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")
TITLE_FONT = Font(bold=True, size=13)
THIN = Side(style="thin", color="BFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TIER_ORDER = {"Primary": 0, "Secondary": 1, "Tertiary": 2}
EARTH_RADIUS_FT = 20902231.0
ILLEGAL_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def log(msg: str) -> None:
    print(msg, file=sys.stderr, flush=True)


# --------------------------------------------------------------------------
# config
# --------------------------------------------------------------------------


def deep_merge(base: dict, override: dict) -> dict:
    out = dict(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = deep_merge(out[key], value)
        else:
            out[key] = value
    return out


def load_config(path: str | None) -> dict:
    if not path:
        return json.loads(json.dumps(DEFAULT_CONFIG))
    text = Path(path).read_text(encoding="utf-8")
    if path.endswith(".json"):
        user = json.loads(text)
    else:
        if yaml is None:
            raise SystemExit("PyYAML is not installed. Run: pip install pyyaml")
        user = yaml.safe_load(text) or {}
    return deep_merge(DEFAULT_CONFIG, user)


# --------------------------------------------------------------------------
# portal session
# --------------------------------------------------------------------------


class Portal:
    """Thin wrapper over the ArcGIS REST endpoints."""

    def __init__(self, portal_url: str, token: str | None) -> None:
        self.portal = portal_url.rstrip("/")
        self.rest = f"{self.portal}/sharing/rest"
        self.token = token
        self.session = requests.Session()

    def request(self, url: str, params: dict, method: str = "GET") -> dict:
        payload = dict(params)
        payload.setdefault("f", "json")
        if self.token:
            payload["token"] = self.token
        if method == "GET":
            resp = self.session.get(url, params=payload, timeout=120)
        else:
            resp = self.session.post(url, data=payload, timeout=180)
        resp.raise_for_status()
        try:
            data = resp.json()
        except ValueError as exc:
            raise RuntimeError(f"{url} did not return JSON: {resp.text[:200]}") from exc
        if isinstance(data, dict) and "error" in data:
            err = data["error"]
            details = "; ".join(err.get("details") or [])
            raise RuntimeError(
                f"ArcGIS error {err.get('code')} on {url}: {err.get('message')} {details}".strip()
            )
        return data


def acquire_token(cfg: dict) -> str | None:
    auth = cfg["auth"]
    mode = (auth.get("mode") or "prompt").lower()
    portal = cfg["portal"].rstrip("/")

    if mode == "anonymous":
        return None

    if mode == "api_key":
        key = os.environ.get(auth.get("api_key_env") or "AGOL_API_KEY")
        if not key:
            raise SystemExit(
                f"auth.mode is api_key but ${auth.get('api_key_env')} is not set."
            )
        return key

    if mode in ("prompt", "password"):
        username = auth.get("username") or os.environ.get("AGOL_USERNAME")
        if not username:
            username = input("AGOL username: ").strip()
        password = os.environ.get("AGOL_PASSWORD") or getpass.getpass(
            f"AGOL password for {username}: "
        )
        resp = requests.post(
            f"{portal}/sharing/rest/generateToken",
            data={
                "username": username,
                "password": password,
                "client": "referer",
                "referer": portal,
                "expiration": 240,
                "f": "json",
            },
            timeout=60,
        )
        resp.raise_for_status()
        data = resp.json()
        if "token" not in data:
            raise SystemExit(f"Sign-in failed: {data}")
        return data["token"]

    raise SystemExit(f"Unknown auth.mode: {mode}")


# --------------------------------------------------------------------------
# layer discovery
# --------------------------------------------------------------------------


class LayerRef:
    def __init__(self, title: str, url: str, group: str = "") -> None:
        self.title = title
        self.url = url.rstrip("/")
        self.group = group
        self.meta: dict[str, Any] = {}

    @property
    def display(self) -> str:
        return f"{self.group} / {self.title}" if self.group else self.title

    def __repr__(self) -> str:  # pragma: no cover - debugging aid
        return f"<LayerRef {self.display} {self.url}>"


_INDEXED_URL = re.compile(r"/\d+$")


def expand_service(portal: Portal, url: str, group: str) -> list[LayerRef]:
    """A FeatureServer/MapServer URL with no layer index -> one ref per sublayer."""
    meta = portal.request(url, {})
    refs: list[LayerRef] = []
    for key in ("layers", "tables"):
        for entry in meta.get(key) or []:
            if entry.get("parentLayerId") not in (None, -1):
                # sublayer of a group layer inside the service; still queryable
                pass
            refs.append(
                LayerRef(entry.get("name") or f"Layer {entry['id']}", f"{url}/{entry['id']}", group)
            )
    return refs


def resolve_entry(portal: Portal, entry: dict, group: str) -> list[LayerRef]:
    title = entry.get("title") or entry.get("name") or entry.get("id") or "Untitled"
    url = entry.get("url")

    if not url and entry.get("itemId"):
        try:
            item = portal.request(f"{portal.rest}/content/items/{entry['itemId']}", {})
            url = item.get("url")
        except RuntimeError as exc:
            log(f"  ! cannot resolve item for '{title}': {exc}")
            return []

    if not url:
        log(f"  - skipping '{title}' (no service URL; basemap or client-side layer)")
        return []

    url = url.rstrip("/")
    if _INDEXED_URL.search(url):
        return [LayerRef(title, url, group)]
    return expand_service(portal, url, group or title)


def layers_from_webmap(portal: Portal, item_id: str) -> list[LayerRef]:
    data = portal.request(f"{portal.rest}/content/items/{item_id}/data", {})
    refs: list[LayerRef] = []

    def walk(entries: Iterable[dict] | None, group: str) -> None:
        for entry in entries or []:
            title = entry.get("title") or "Group"
            if entry.get("layerType") == "GroupLayer" or (
                entry.get("layers") and not entry.get("url") and not entry.get("itemId")
            ):
                walk(entry.get("layers"), f"{group} / {title}" if group else title)
                continue
            refs.extend(resolve_entry(portal, entry, group))

    walk(data.get("operationalLayers"), "")
    walk(data.get("tables"), "Tables")
    return refs


def discover_layers(portal: Portal, cfg: dict) -> list[LayerRef]:
    src = cfg["source"]
    refs: list[LayerRef] = []
    if src.get("webmap_item_id"):
        refs.extend(layers_from_webmap(portal, src["webmap_item_id"]))
    for url in src.get("service_urls") or []:
        url = url.rstrip("/")
        if _INDEXED_URL.search(url):
            meta = portal.request(url, {})
            refs.append(LayerRef(meta.get("name") or url.rsplit("/", 2)[-2], url, ""))
        else:
            refs.extend(expand_service(portal, url, ""))
    if not refs:
        raise SystemExit(
            "No layers found. Set source.webmap_item_id or source.service_urls in the config."
        )
    return filter_layers(refs, cfg)


def _matches_any(text: str, needles: Sequence[str]) -> bool:
    lowered = text.lower()
    return any(n.lower() in lowered for n in needles if n)


def filter_layers(refs: list[LayerRef], cfg: dict) -> list[LayerRef]:
    only = cfg["query"].get("only_layers") or []
    skip = cfg["query"].get("skip_layers") or []
    out = []
    for ref in refs:
        if only and not _matches_any(ref.display, only):
            continue
        if skip and _matches_any(ref.display, skip):
            log(f"  - skipping '{ref.display}' (query.skip_layers)")
            continue
        out.append(ref)
    return out


# --------------------------------------------------------------------------
# feature querying
# --------------------------------------------------------------------------


def load_metadata(portal: Portal, ref: LayerRef) -> dict:
    if not ref.meta:
        ref.meta = portal.request(ref.url, {})
    return ref.meta


def field_map(meta: dict) -> dict[str, dict]:
    return {f["name"]: f for f in meta.get("fields") or []}


def build_domain_lookup(meta: dict) -> dict[str, dict[Any, str]]:
    """field name -> {code: description} for coded-value domains."""
    lookup: dict[str, dict[Any, str]] = {}
    for field in meta.get("fields") or []:
        domain = field.get("domain") or {}
        if domain.get("type") == "codedValue":
            lookup[field["name"]] = {
                cv["code"]: cv["name"] for cv in domain.get("codedValues") or []
            }
    # subtype-driven domains: use the subtype field's own coded values
    subtype_field = meta.get("typeIdField")
    if subtype_field:
        codes = {}
        for stype in meta.get("types") or []:
            codes[stype.get("id")] = stype.get("name")
        if codes:
            lookup.setdefault(subtype_field, {}).update(codes)
    return lookup


def query_features(portal: Portal, ref: LayerRef, cfg: dict) -> list[dict]:
    meta = load_metadata(portal, ref)
    q = cfg["query"]
    page_size = min(int(q["page_size"]), int(meta.get("maxRecordCount") or 1000))
    oid_field = meta.get("objectIdField") or "OBJECTID"
    supports_paging = bool(
        (meta.get("advancedQueryCapabilities") or {}).get("supportsPagination")
    )
    limit = int(q["max_features_per_layer"])

    base = {
        "where": q["where"],
        "outFields": "*",
        "returnGeometry": True,
        "outSR": q["out_sr"],
        "f": "json",
    }

    features: list[dict] = []
    if supports_paging:
        offset = 0
        while True:
            params = dict(base, resultOffset=offset, resultRecordCount=page_size)
            data = portal.request(f"{ref.url}/query", params, method="POST")
            batch = data.get("features") or []
            features.extend(batch)
            if len(batch) < page_size or len(features) >= limit:
                break
            offset += len(batch)
    else:
        last_oid = -1
        while True:
            where = f"({q['where']}) AND {oid_field} > {last_oid}"
            params = dict(base, where=where, orderByFields=oid_field, resultRecordCount=page_size)
            data = portal.request(f"{ref.url}/query", params, method="POST")
            batch = data.get("features") or []
            if not batch:
                break
            features.extend(batch)
            last_oid = batch[-1]["attributes"].get(oid_field, last_oid)
            if len(batch) < page_size or len(features) >= limit:
                break

    if len(features) > limit:
        log(f"  ! '{ref.display}' truncated at max_features_per_layer={limit}")
        features = features[:limit]
    return features


# --------------------------------------------------------------------------
# geometry + value coercion
# --------------------------------------------------------------------------


def representative_point(geometry: dict | None) -> tuple[float | None, float | None]:
    """Return (lon, lat) for any geometry type; centroid for lines/polygons."""
    if not geometry:
        return None, None
    if "x" in geometry and "y" in geometry:
        x, y = geometry.get("x"), geometry.get("y")
        return (x, y) if isinstance(x, (int, float)) and isinstance(y, (int, float)) else (None, None)
    if geometry.get("points"):
        pts = geometry["points"]
    elif geometry.get("paths"):
        pts = [p for path in geometry["paths"] for p in path]
    elif geometry.get("rings"):
        pts = [p for ring in geometry["rings"] for p in ring]
    else:
        return None, None
    pts = [p for p in pts if len(p) >= 2 and isinstance(p[0], (int, float))]
    if not pts:
        return None, None
    return sum(p[0] for p in pts) / len(pts), sum(p[1] for p in pts) / len(pts)


def coerce(value: Any, field: dict | None, domains: dict[str, dict]) -> Any:
    if value is None:
        return None
    name = (field or {}).get("name")
    if name and name in domains and value in domains[name]:
        return domains[name][value]
    ftype = (field or {}).get("type")
    if ftype in ("esriFieldTypeDate",) and isinstance(value, (int, float)):
        try:
            return datetime.fromtimestamp(value / 1000, tz=timezone.utc).replace(tzinfo=None)
        except (OverflowError, OSError, ValueError):
            return value
    if isinstance(value, (dict, list)):
        return json.dumps(value)
    return value


def haversine_ft(lon1: float, lat1: float, lon2: float, lat2: float) -> float:
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = p2 - p1
    dlam = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlam / 2) ** 2
    return 2 * EARTH_RADIUS_FT * math.asin(math.sqrt(a))


# --------------------------------------------------------------------------
# tier classification
# --------------------------------------------------------------------------


def pick_field(attrs: dict, candidates: Sequence[str]) -> tuple[str | None, Any]:
    """Case-insensitive first-hit lookup across candidate field names."""
    lowered = {k.lower(): k for k in attrs}
    for cand in candidates or []:
        key = lowered.get(cand.lower())
        if key is not None and attrs[key] not in (None, "", " "):
            return key, attrs[key]
    return None, None


def classify_tier(attrs: dict, rules: dict) -> tuple[str, str]:
    """Return (tier, how_it_was_determined)."""
    field, value = pick_field(attrs, rules.get("fields") or [])
    if value is not None:
        text = str(value).strip().lower()
        for tier, keywords in (rules.get("keywords") or {}).items():
            for kw in keywords:
                kw = str(kw).strip().lower()
                if not kw:
                    continue
                if text == kw or re.search(rf"(?<![a-z0-9]){re.escape(kw)}(?![a-z0-9])", text):
                    return tier, f"{field}={value}"

    ratio_map = {str(k).lower(): v for k, v in (rules.get("ratio_map") or {}).items()}
    if ratio_map:
        rfield, rvalue = pick_field(attrs, rules.get("ratio_fields") or [])
        if rvalue is not None:
            key = str(rvalue).strip().lower().replace(":", "x").replace(" ", "")
            if key in ratio_map:
                return ratio_map[key], f"{rfield}={rvalue}"

    for pattern, tier in (rules.get("name_patterns") or {}).items():
        for attr_value in attrs.values():
            if attr_value is None:
                continue
            if re.search(pattern, str(attr_value), re.IGNORECASE):
                return tier, f"pattern /{pattern}/"

    return "", "unresolved"


# --------------------------------------------------------------------------
# address resolution
# --------------------------------------------------------------------------


class AddressIndex:
    """Grid-bucketed nearest-neighbour lookup over an address point layer."""

    CELL = 0.002  # ~700 ft of latitude

    def __init__(self) -> None:
        self.cells: dict[tuple[int, int], list[tuple[float, float, str]]] = defaultdict(list)
        self.count = 0

    def add(self, lon: float, lat: float, label: str) -> None:
        self.cells[(int(lat / self.CELL), int(lon / self.CELL))].append((lon, lat, label))
        self.count += 1

    def nearest(self, lon: float, lat: float, max_ft: float) -> tuple[str, float | None]:
        """Nearest labelled point, searching outward one ring of cells at a time."""
        if not self.count:
            return "", None
        # Feet per cell, using the shorter of the lat/lon conversions so the
        # early-exit bound below is always conservative.
        cell_ft = self.CELL * 364000.0 * max(math.cos(math.radians(lat)), 0.05)
        max_rings = int(max_ft / cell_ft) + 2
        row, col = int(lat / self.CELL), int(lon / self.CELL)
        best_label, best_dist = "", None

        for radius in range(max_rings + 1):
            for dr in range(-radius, radius + 1):
                for dc in range(-radius, radius + 1):
                    if radius and max(abs(dr), abs(dc)) != radius:
                        continue  # inner cells were scanned on an earlier pass
                    for clon, clat, label in self.cells.get((row + dr, col + dc), ()):
                        dist = haversine_ft(lon, lat, clon, clat)
                        if best_dist is None or dist < best_dist:
                            best_label, best_dist = label, dist
            # Anything still unscanned is at least `radius` cells away, so once
            # the best hit is closer than that we can stop.
            if best_dist is not None and best_dist <= radius * cell_ft:
                break

        if best_dist is not None and best_dist > max_ft:
            return "", best_dist
        return best_label, best_dist


def compose_address(attrs: dict, fields: Sequence[str]) -> str:
    lowered = {k.lower(): k for k in attrs}
    parts = []
    for name in fields:
        key = lowered.get(str(name).lower())
        if key and attrs[key] not in (None, "", " "):
            parts.append(str(attrs[key]).strip())
    return " ".join(parts).strip()


def build_address_index(tables: dict[str, "LayerTable"], cfg: dict) -> AddressIndex:
    index = AddressIndex()
    needle = cfg["address"].get("layer") or ""
    fields = cfg["address"].get("fields") or []
    for table in tables.values():
        if not needle or not _matches_any(table.ref.display, [needle]):
            continue
        for row in table.rows:
            lon, lat = row["_lon"], row["_lat"]
            label = compose_address(row["_raw"], fields)
            if lon is None or lat is None or not label:
                continue
            index.add(lon, lat, label)
        log(f"  address source: '{table.ref.display}' ({index.count} points)")
    if not index.count:
        log(f"  ! no address points found matching address.layer='{needle}'")
    return index


class Geocoder:
    def __init__(self, portal: Portal, cfg: dict) -> None:
        self.portal = portal
        self.url = cfg["address"]["geocode_url"]
        self.cache_path = Path(cfg["address"].get("geocode_cache") or "geocode_cache.json")
        self.cache: dict[str, str] = {}
        if self.cache_path.exists():
            try:
                self.cache = json.loads(self.cache_path.read_text(encoding="utf-8"))
            except ValueError:
                self.cache = {}
        self.calls = 0

    def reverse(self, lon: float, lat: float) -> str:
        key = f"{lon:.6f},{lat:.6f}"
        if key in self.cache:
            return self.cache[key]
        try:
            data = self.portal.request(
                self.url,
                {"location": json.dumps({"x": lon, "y": lat, "spatialReference": {"wkid": 4326}}),
                 "outSR": 4326, "featureTypes": "PointAddress,StreetAddress"},
            )
            label = (data.get("address") or {}).get("LongLabel") or (
                data.get("address") or {}
            ).get("Match_addr") or ""
        except (RuntimeError, requests.RequestException) as exc:
            log(f"  ! reverse geocode failed at {key}: {exc}")
            label = ""
        self.calls += 1
        self.cache[key] = label
        return label

    def save(self) -> None:
        if self.calls:
            self.cache_path.write_text(json.dumps(self.cache, indent=1), encoding="utf-8")
            log(f"  reverse geocoded {self.calls} new point(s); cache -> {self.cache_path}")


# --------------------------------------------------------------------------
# in-memory tables
# --------------------------------------------------------------------------


class LayerTable:
    def __init__(self, ref: LayerRef, columns: list[str], rows: list[dict]) -> None:
        self.ref = ref
        self.columns = columns
        self.rows = rows


def build_table(portal: Portal, ref: LayerRef, cfg: dict) -> LayerTable:
    meta = load_metadata(portal, ref)
    fields = field_map(meta)
    domains = build_domain_lookup(meta) if cfg["query"]["decode_domains"] else {}
    features = query_features(portal, ref, cfg)

    columns: list[str] = [f["name"] for f in meta.get("fields") or []]
    rows: list[dict] = []
    for feat in features:
        attrs = feat.get("attributes") or {}
        for key in attrs:
            if key not in columns:
                columns.append(key)
        lon, lat = representative_point(feat.get("geometry"))
        row = {
            "_raw": attrs,
            "_lon": lon,
            "_lat": lat,
            "_values": {k: coerce(v, fields.get(k), domains) for k, v in attrs.items()},
        }
        rows.append(row)
    return LayerTable(ref, columns, rows)


# --------------------------------------------------------------------------
# splitter sheet
# --------------------------------------------------------------------------


SPLITTER_COLUMNS = [
    "Splitter",
    "Tier",
    "Ratio",
    "Nearest Address",
    "Distance (ft)",
    "Latitude",
    "Longitude",
    "Layer",
    "Tier Source",
    "QA Flags",
]


def build_splitter_rows(
    tables: dict[str, LayerTable], cfg: dict, index: AddressIndex, geocoder: Geocoder | None
) -> list[dict]:
    scfg = cfg["splitters"]
    acfg = cfg["address"]
    needles = scfg.get("layers") or []
    max_ft = float(acfg.get("max_distance_ft") or 500)
    source = (acfg.get("source") or "none").lower()

    out: list[dict] = []
    seen_names: Counter[str] = Counter()

    for table in tables.values():
        if needles and not _matches_any(table.ref.display, needles):
            continue
        for row in table.rows:
            attrs = row["_values"]
            _, name = pick_field(attrs, scfg.get("name_fields") or [])
            _, ratio = pick_field(attrs, scfg.get("ratio_fields") or [])
            tier_rules = dict(scfg.get("tier") or {})
            tier_rules.setdefault("ratio_fields", scfg.get("ratio_fields") or [])
            tier, how = classify_tier(attrs, tier_rules)

            lon, lat = row["_lon"], row["_lat"]
            address, distance = "", None
            if lon is not None and lat is not None:
                if source == "layer":
                    address, distance = index.nearest(lon, lat, max_ft)
                elif source == "geocode" and geocoder is not None:
                    address = geocoder.reverse(lon, lat)

            flags = []
            if not tier:
                flags.append("no tier")
            if not address:
                flags.append("no address")
            elif distance is not None and distance > max_ft:
                flags.append(f"address >{int(max_ft)} ft")
            if not ratio:
                flags.append("no ratio")
            if lon is None or lat is None:
                flags.append("no geometry")

            label = str(name) if name is not None else ""
            seen_names[label] += 1

            out.append(
                {
                    "Splitter": label,
                    "Tier": tier,
                    "Ratio": ratio,
                    "Nearest Address": address,
                    "Distance (ft)": round(distance, 1) if distance is not None else None,
                    "Latitude": round(lat, 6) if lat is not None else None,
                    "Longitude": round(lon, 6) if lon is not None else None,
                    "Layer": table.ref.display,
                    "Tier Source": how,
                    "QA Flags": ", ".join(flags),
                    "_attrs": attrs,
                }
            )

    for row in out:
        if row["Splitter"] and seen_names[row["Splitter"]] > 1:
            row["QA Flags"] = ", ".join(filter(None, [row["QA Flags"], "duplicate name"]))

    out.sort(key=lambda r: (TIER_ORDER.get(r["Tier"], 9), str(r["Splitter"])))
    return out


# --------------------------------------------------------------------------
# workbook writing
# --------------------------------------------------------------------------


def safe_sheet_name(name: str, used: set[str]) -> str:
    clean = ILLEGAL_SHEET_CHARS.sub("-", name).strip() or "Layer"
    clean = clean[:31]
    base, n = clean, 2
    while clean.lower() in used:
        suffix = f"~{n}"
        clean = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(clean.lower())
    return clean


def style_sheet(ws, header_row: int, ncols: int, nrows: int, max_width: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if nrows:
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(ncols)}{header_row + nrows}"
        )
    for col in range(1, ncols + 1):
        letter = get_column_letter(col)
        widest = 0
        for row in range(header_row, min(header_row + nrows, header_row + 400) + 1):
            value = ws.cell(row=row, column=col).value
            if value is not None:
                widest = max(widest, len(str(value)))
        ws.column_dimensions[letter].width = max(9, min(max_width, widest + 2))


def write_layer_sheet(wb: Workbook, table: LayerTable, used: set[str], cfg: dict) -> str:
    name = safe_sheet_name(table.ref.title, used)
    ws = wb.create_sheet(name)

    ws.cell(row=1, column=1, value=table.ref.display).font = TITLE_FONT
    ws.cell(row=2, column=1, value=table.ref.url).font = Font(color="808080", size=9)

    has_geom = any(r["_lon"] is not None for r in table.rows)
    columns = list(table.columns) + (["Longitude", "Latitude"] if has_geom else [])

    header_row = 4
    for idx, col in enumerate(columns, start=1):
        ws.cell(row=header_row, column=idx, value=col)
    for r, row in enumerate(table.rows, start=header_row + 1):
        for c, col in enumerate(columns, start=1):
            if col == "Longitude":
                value = round(row["_lon"], 6) if row["_lon"] is not None else None
            elif col == "Latitude":
                value = round(row["_lat"], 6) if row["_lat"] is not None else None
            else:
                value = row["_values"].get(col)
            ws.cell(row=r, column=c, value=value)

    style_sheet(ws, header_row, len(columns), len(table.rows), cfg["output"]["max_column_width"])
    return name


def write_splitter_sheet(wb: Workbook, rows: list[dict], cfg: dict) -> None:
    ws = wb.create_sheet("Splitters", 0)
    ws.cell(row=1, column=1, value="DFN Splitter Schedule").font = Font(bold=True, size=15)
    ws.cell(
        row=2,
        column=1,
        value=f"Generated {datetime.now():%Y-%m-%d %H:%M} | {len(rows)} splitters | "
        f"{sum(1 for r in rows if r['QA Flags'])} needing review",
    ).font = Font(color="808080", size=9)

    header_row = 4
    for idx, col in enumerate(SPLITTER_COLUMNS, start=1):
        ws.cell(row=header_row, column=idx, value=col)
    for r, row in enumerate(rows, start=header_row + 1):
        for c, col in enumerate(SPLITTER_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(col))
            cell.border = CELL_BORDER
            if row["QA Flags"]:
                cell.fill = FLAG_FILL

    style_sheet(ws, header_row, len(SPLITTER_COLUMNS), len(rows), cfg["output"]["max_column_width"])


def write_index_sheet(wb: Workbook, entries: list[tuple[str, str, int]]) -> None:
    ws = wb.create_sheet("Index", 0)
    ws.cell(row=1, column=1, value="Layers in this export").font = Font(bold=True, size=15)
    header_row = 3
    for idx, col in enumerate(["Sheet", "Layer", "Features"], start=1):
        ws.cell(row=header_row, column=idx, value=col)
    for r, (sheet, display, count) in enumerate(entries, start=header_row + 1):
        cell = ws.cell(row=r, column=1, value=sheet)
        cell.hyperlink = f"#'{sheet}'!A1"
        cell.font = Font(color="0563C1", underline="single")
        ws.cell(row=r, column=2, value=display)
        ws.cell(row=r, column=3, value=count)
    style_sheet(ws, header_row, 3, len(entries), 60)


def write_csvs(tables: dict[str, LayerTable], directory: str) -> None:
    out_dir = Path(directory)
    out_dir.mkdir(parents=True, exist_ok=True)
    for table in tables.values():
        name = re.sub(r"[^A-Za-z0-9_.-]+", "_", table.ref.title)[:80] or "layer"
        path = out_dir / f"{name}.csv"
        columns = list(table.columns) + ["Longitude", "Latitude"]
        with path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=columns)
            writer.writeheader()
            for row in table.rows:
                record = {k: row["_values"].get(k) for k in table.columns}
                record["Longitude"] = row["_lon"]
                record["Latitude"] = row["_lat"]
                writer.writerow(record)
    log(f"  csv -> {out_dir}")


# --------------------------------------------------------------------------
# commands
# --------------------------------------------------------------------------


def cmd_discover(portal: Portal, cfg: dict, args) -> int:
    refs = discover_layers(portal, cfg)
    print(f"\n{len(refs)} layer(s) found\n" + "=" * 72)
    for ref in refs:
        meta = load_metadata(portal, ref)
        try:
            count = portal.request(
                f"{ref.url}/query", {"where": cfg["query"]["where"], "returnCountOnly": True}
            ).get("count", "?")
        except RuntimeError as exc:
            count = f"error: {exc}"
        print(f"\n{ref.display}")
        print(f"  url        {ref.url}")
        print(f"  geometry   {meta.get('geometryType') or 'table (no geometry)'}")
        print(f"  features   {count}")
        print("  fields:")
        for field in meta.get("fields") or []:
            domain = field.get("domain") or {}
            tag = "  [coded domain]" if domain.get("type") == "codedValue" else ""
            print(
                f"    {field['name']:<34} {field.get('type','').replace('esriFieldType',''):<12}"
                f" {field.get('alias','')}{tag}"
            )
    print("\nCopy the field names you want into config.yaml, then run: profile, then export.")
    return 0


def cmd_profile(portal: Portal, cfg: dict, args) -> int:
    refs = discover_layers(portal, cfg)
    needles = args.layer or cfg["splitters"]["layers"]
    targets = [r for r in refs if not needles or _matches_any(r.display, needles)]
    if not targets:
        raise SystemExit(f"No layer matched {needles}")
    for ref in targets:
        table = build_table(portal, ref, cfg)
        print(f"\n{ref.display}  ({len(table.rows)} features)\n" + "=" * 72)
        for col in table.columns:
            values = Counter(
                str(r["_values"].get(col)) for r in table.rows if r["_values"].get(col) not in (None, "")
            )
            if not values or len(values) > int(args.max_distinct):
                continue
            rendered = ", ".join(f"{v} ({n})" for v, n in values.most_common(int(args.max_distinct)))
            print(f"  {col}: {rendered}")
    print("\nUse these values to fill splitters.tier.keywords in the config.")
    return 0


def cmd_export(portal: Portal, cfg: dict, args) -> int:
    refs = discover_layers(portal, cfg)
    log(f"exporting {len(refs)} layer(s)")

    tables: dict[str, LayerTable] = {}
    for ref in refs:
        try:
            table = build_table(portal, ref, cfg)
        except (RuntimeError, requests.RequestException) as exc:
            log(f"  ! '{ref.display}' failed: {exc}")
            continue
        tables[ref.url] = table
        log(f"  {ref.display}: {len(table.rows)} features")

    if not tables:
        raise SystemExit("Nothing exported -- every layer failed or was filtered out.")

    source = (cfg["address"].get("source") or "none").lower()
    index = build_address_index(tables, cfg) if source == "layer" else AddressIndex()
    geocoder = Geocoder(portal, cfg) if source == "geocode" else None

    splitter_rows = build_splitter_rows(tables, cfg, index, geocoder)
    if geocoder:
        geocoder.save()

    wb = Workbook()
    wb.remove(wb.active)

    entries: list[tuple[str, str, int]] = []
    used: set[str] = {"index", "splitters"}
    for table in tables.values():
        sheet = write_layer_sheet(wb, table, used, cfg)
        entries.append((sheet, table.ref.display, len(table.rows)))

    if splitter_rows:
        write_splitter_sheet(wb, splitter_rows, cfg)
        log(
            f"  splitters: {len(splitter_rows)} rows, "
            f"{sum(1 for r in splitter_rows if r['QA Flags'])} flagged for review"
        )
    else:
        log("  ! no splitter rows -- check splitters.layers in the config")

    write_index_sheet(wb, entries)

    out_path = Path(args.out or cfg["output"]["workbook"])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log(f"\nworkbook -> {out_path.resolve()}")

    if cfg["output"].get("csv_dir"):
        write_csvs(tables, cfg["output"]["csv_dir"])
    return 0


# --------------------------------------------------------------------------
# entry point
# --------------------------------------------------------------------------


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Export ArcGIS Online layers to Excel, with a DFN splitter schedule."
    )
    parser.add_argument("-c", "--config", help="path to config.yaml")
    parser.add_argument("--portal", help="override portal URL")
    parser.add_argument("--webmap", help="override source.webmap_item_id")
    parser.add_argument("--anonymous", action="store_true", help="skip sign-in (public data)")

    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("discover", help="list layers and fields")

    p_profile = sub.add_parser("profile", help="show distinct field values")
    p_profile.add_argument("--layer", action="append", help="layer name substring (repeatable)")
    p_profile.add_argument("--max-distinct", default=25, help="hide fields with more values")

    p_export = sub.add_parser("export", help="write the workbook")
    p_export.add_argument("-o", "--out", help="output .xlsx path")

    args = parser.parse_args(argv)
    cfg = load_config(args.config)
    if args.portal:
        cfg["portal"] = args.portal
    if args.webmap:
        cfg["source"]["webmap_item_id"] = args.webmap
    if args.anonymous:
        cfg["auth"]["mode"] = "anonymous"

    token = acquire_token(cfg)
    portal = Portal(cfg["portal"], token)

    handlers = {"discover": cmd_discover, "profile": cmd_profile, "export": cmd_export}
    try:
        return handlers[args.command](portal, cfg, args)
    except requests.HTTPError as exc:
        raise SystemExit(f"HTTP error: {exc}") from exc


if __name__ == "__main__":
    sys.exit(main())
