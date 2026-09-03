#!/usr/bin/env python3
"""
Offline test for dfn_export.py.

Stands up a fake ArcGIS REST service in memory -- a web map holding a splitter
layer, an address point layer and a fiber cable line layer -- runs a full
export against it, and checks the workbook that comes out.

    python3 test_dfn_export.py
"""

import json
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent))

import dfn_export as dx
from openpyxl import load_workbook

SVC = "https://services9.arcgis.com/FAKE/arcgis/rest/services/DFN/FeatureServer"

SPLITTER_FIELDS = [
    {"name": "OBJECTID", "type": "esriFieldTypeOID", "alias": "OBJECTID"},
    {"name": "SPLITTER_ID", "type": "esriFieldTypeString", "alias": "Splitter ID"},
    {
        "name": "SPLTR_LEVEL",
        "type": "esriFieldTypeSmallInteger",
        "alias": "Level",
        "domain": {
            "type": "codedValue",
            "codedValues": [
                {"code": 1, "name": "Level 1"},
                {"code": 2, "name": "Level 2"},
                {"code": 3, "name": "Level 3"},
            ],
        },
    },
    {"name": "SPLIT_RATIO", "type": "esriFieldTypeString", "alias": "Ratio"},
    {"name": "INSTALLED", "type": "esriFieldTypeDate", "alias": "Installed"},
]

ADDRESS_FIELDS = [
    {"name": "OBJECTID", "type": "esriFieldTypeOID", "alias": "OBJECTID"},
    {"name": "HOUSE_NUM", "type": "esriFieldTypeString", "alias": "House"},
    {"name": "STREET", "type": "esriFieldTypeString", "alias": "Street"},
]

CABLE_FIELDS = [
    {"name": "OBJECTID", "type": "esriFieldTypeOID", "alias": "OBJECTID"},
    {"name": "CABLE_NAME", "type": "esriFieldTypeString", "alias": "Cable"},
]


def point(x, y):
    return {"x": x, "y": y}


SPLITTERS = [
    # id, level code, ratio, lon, lat
    ("SPL-001", 1, "1x8", -83.0000, 40.0000),
    ("SPL-002", 2, "1x16", -83.0010, 40.0005),
    ("SPL-003", 3, "1x32", -83.0020, 40.0010),
    ("SPL-004", None, "1x32", -83.0500, 40.0500),  # no tier, no nearby address
]

ADDRESSES = [
    ("100", "N HIGH ST", -83.00005, 40.00005),
    ("212", "W 5TH AVE", -83.00105, 40.00055),
    ("415", "E LANE AVE", -83.00205, 40.00105),
]


class FakePortal(dx.Portal):
    def __init__(self):
        super().__init__("https://fake.maps.arcgis.com", "TOKEN")
        self.calls = []

    def request(self, url, params, method="GET"):
        self.calls.append((url, params.get("where"), method))

        if url.endswith("/data"):
            return {
                "operationalLayers": [
                    {
                        "layerType": "GroupLayer",
                        "title": "DFN 1234",
                        "layers": [
                            {"title": "Splitters", "url": f"{SVC}/0"},
                            {"title": "Fiber Cable", "url": f"{SVC}/2"},
                        ],
                    },
                    {"title": "Address Points", "url": f"{SVC}/1"},
                    {"title": "Sketch layer", "layerType": "ArcGISFeatureLayer"},  # no url
                ]
            }

        if url == f"{SVC}/0":
            return {
                "name": "Splitters",
                "geometryType": "esriGeometryPoint",
                "objectIdField": "OBJECTID",
                "maxRecordCount": 2,  # forces paging
                "advancedQueryCapabilities": {"supportsPagination": True},
                "fields": SPLITTER_FIELDS,
            }
        if url == f"{SVC}/1":
            return {
                "name": "Address Points",
                "geometryType": "esriGeometryPoint",
                "objectIdField": "OBJECTID",
                "maxRecordCount": 1000,
                "advancedQueryCapabilities": {"supportsPagination": True},
                "fields": ADDRESS_FIELDS,
            }
        if url == f"{SVC}/2":
            return {
                "name": "Fiber Cable",
                "geometryType": "esriGeometryPolyline",
                "objectIdField": "OBJECTID",
                "maxRecordCount": 1000,
                "advancedQueryCapabilities": {"supportsPagination": True},
                "fields": CABLE_FIELDS,
            }

        if url == f"{SVC}/0/query":
            feats = [
                {
                    "attributes": {
                        "OBJECTID": i + 1,
                        "SPLITTER_ID": sid,
                        "SPLTR_LEVEL": lvl,
                        "SPLIT_RATIO": ratio,
                        "INSTALLED": 1700000000000,
                    },
                    "geometry": point(lon, lat),
                }
                for i, (sid, lvl, ratio, lon, lat) in enumerate(SPLITTERS)
            ]
            return {"features": self._page(feats, params)}

        if url == f"{SVC}/1/query":
            feats = [
                {
                    "attributes": {"OBJECTID": i + 1, "HOUSE_NUM": num, "STREET": street},
                    "geometry": point(lon, lat),
                }
                for i, (num, street, lon, lat) in enumerate(ADDRESSES)
            ]
            return {"features": self._page(feats, params)}

        if url == f"{SVC}/2/query":
            feats = [
                {
                    "attributes": {"OBJECTID": 1, "CABLE_NAME": "F-001"},
                    "geometry": {"paths": [[[-83.0, 40.0], [-83.002, 40.001]]]},
                }
            ]
            return {"features": self._page(feats, params)}

        raise AssertionError(f"unexpected request to {url}")

    @staticmethod
    def _page(feats, params):
        offset = int(params.get("resultOffset") or 0)
        size = int(params.get("resultRecordCount") or len(feats))
        return feats[offset : offset + size]


CONFIG = dx.deep_merge(
    dx.DEFAULT_CONFIG,
    {
        "source": {"webmap_item_id": "abc123"},
        "splitters": {
            "layers": ["splitter"],
            "name_fields": ["SPLITTER_ID"],
            "ratio_fields": ["SPLIT_RATIO"],
            "tier": {
                "fields": ["SPLTR_LEVEL"],
                "keywords": {
                    "Primary": ["level 1"],
                    "Secondary": ["level 2"],
                    "Tertiary": ["level 3"],
                },
            },
        },
        "address": {
            "source": "layer",
            "layer": "address",
            "fields": ["HOUSE_NUM", "STREET"],
            "max_distance_ft": 300,
        },
    },
)


def check(condition, message):
    if not condition:
        raise AssertionError(message)
    print(f"  ok  {message}")


def main() -> int:
    portal = FakePortal()
    tmp = Path(tempfile.mkdtemp())
    out = tmp / "dfn_export.xlsx"

    print("running export against fake service...")
    rc = dx.cmd_export(portal, CONFIG, SimpleNamespace(out=str(out)))
    check(rc == 0, "export returned 0")
    check(out.exists(), "workbook was written")

    wb = load_workbook(out)
    check(wb.sheetnames[0] == "Index", "Index is the first sheet")
    check("Splitters" in wb.sheetnames, "Splitters sheet exists")
    check("Fiber Cable" in wb.sheetnames, "per-layer sheet exists for a line layer")

    ws = wb["Splitters"]
    header = [c.value for c in ws[4]]
    check(header == dx.SPLITTER_COLUMNS, "splitter header matches spec")

    rows = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        rows[r[0]] = dict(zip(header, r))
    check(len(rows) == 4, f"all 4 splitters exported (got {len(rows)})")

    check(rows["SPL-001"]["Tier"] == "Primary", "coded domain 1 -> Primary")
    check(rows["SPL-002"]["Tier"] == "Secondary", "coded domain 2 -> Secondary")
    check(rows["SPL-003"]["Tier"] == "Tertiary", "coded domain 3 -> Tertiary")
    check(not rows["SPL-004"]["Tier"], "unmapped level leaves tier blank")

    check(
        rows["SPL-001"]["Nearest Address"] == "100 N HIGH ST",
        f"nearest address resolved ({rows['SPL-001']['Nearest Address']})",
    )
    check(
        rows["SPL-002"]["Nearest Address"] == "212 W 5TH AVE",
        "second splitter matched its own address, not the first",
    )
    check(rows["SPL-001"]["Distance (ft)"] < 30, "distance is in feet and small")
    check("no tier" in rows["SPL-004"]["QA Flags"], "missing tier is flagged")
    check("no address" in rows["SPL-004"]["QA Flags"], "far-away splitter flagged as no address")
    check(not rows["SPL-001"]["QA Flags"], "clean row carries no flags")

    tiers = [rows[k]["Tier"] for k in ["SPL-001", "SPL-002", "SPL-003"]]
    ordered = [c.value for c in ws["B"][4:8]]
    check(ordered[:3] == tiers, "rows sorted primary -> secondary -> tertiary")

    ws_spl = wb["Splitters"]
    check(ws_spl.freeze_panes == "A5", "header row frozen")
    check(ws_spl.auto_filter.ref is not None, "autofilter applied")

    ws_layer = wb["Splitters" if "Splitters" not in wb.sheetnames else "Fiber Cable"]
    check(
        "Longitude" in [c.value for c in ws_layer[4]],
        "line layer got a representative lon/lat column",
    )

    # paging: maxRecordCount 2 over 4 splitters must mean more than one query call
    splitter_queries = [c for c in portal.calls if c[0].endswith("/0/query")]
    check(len(splitter_queries) >= 2, f"paging issued {len(splitter_queries)} requests")

    print("\nall checks passed")
    print(f"sample workbook: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
